from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from formatters.utils import sanitize_text

def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = "LegalEase - AI-Powered Legal Document Generator"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for raw in sanitize_text(text).split("\n"):
        line = raw.strip()
        if not line:
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        if (line.isupper() and len(line) < 100) or line.endswith(":"):
            r = p.add_run(line)
            r.bold = True
        else:
            r = p.add_run(line)

        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
